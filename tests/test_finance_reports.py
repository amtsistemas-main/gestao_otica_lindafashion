from datetime import date

import pytest

from modules.finance import reports

TODAY = date(2026, 10, 26)  # segunda-feira

ENTRIES = [
    {"id": "1", "direction": "entrada", "category": "Vendas", "description": "Venda V-1", "amount": 500,
     "due_date": "2026-10-20", "status": "pago", "paid_at": "2026-10-21"},
    {"id": "2", "direction": "saida", "category": "Aluguel", "description": "Aluguel", "amount": 2000,
     "due_date": "2026-10-30", "status": "pendente", "paid_at": None},
    {"id": "3", "direction": "entrada", "category": "Vendas", "description": "Parcela 2/3", "amount": 150,
     "due_date": "2026-10-10", "status": "pendente", "paid_at": None},
    {"id": "4", "direction": "saida", "category": "Fornecedor", "description": "Cancelado", "amount": 999,
     "due_date": "2026-10-15", "status": "cancelado", "paid_at": None},
    {"id": "5", "direction": "saida", "category": "Energia", "description": "Luz", "amount": 300,
     "due_date": "2026-09-28", "status": "pago", "paid_at": "2026-10-02"},
]


def test_periods_cover_whole_calendar_units():
    assert reports.period_bounds("Hoje", TODAY) == (TODAY, TODAY)
    assert reports.period_bounds("Esta semana", TODAY) == (date(2026, 10, 25), date(2026, 10, 31))
    # O mês vai até o último dia: contas que vencem no fim do mês entram no filtro.
    assert reports.period_bounds("Este mês", TODAY) == (date(2026, 10, 1), date(2026, 10, 31))
    assert reports.period_bounds("Este mês", date(2026, 2, 10)) == (date(2026, 2, 1), date(2026, 2, 28))
    assert reports.period_bounds("Este ano", TODAY) == (date(2026, 1, 1), date(2026, 12, 31))
    assert reports.period_bounds("Personalizado", TODAY, date(2026, 10, 31), date(2026, 10, 1)) == (date(2026, 10, 1), date(2026, 10, 31))
    assert reports.period_bounds("Todo o período", TODAY) == (None, None)


def test_paid_entries_count_on_payment_day():
    luz = ENTRIES[4]
    assert reports.entry_day(luz) == date(2026, 10, 2)
    assert reports.in_period(luz, date(2026, 10, 1), date(2026, 10, 31))
    assert not reports.in_period(luz, date(2026, 9, 1), date(2026, 9, 30))


def test_report_totals_and_categories():
    start, end = reports.period_bounds("Este mês", TODAY)
    report = reports.build_report(ENTRIES, "Ótica Teste", "Este mês", start, end, today=TODAY)
    assert [row["id"] for row in report.rows] == ["5", "3", "1", "2"]  # ordem cronológica, sem cancelados
    assert report.income_paid == 500 and report.expense_paid == 300 and report.balance == 200
    assert report.receivable_open == 150 and report.payable_open == 2000
    assert report.overdue == 150  # parcela vencida em 10/10
    assert report.cancelled == 1
    assert report.by_category[0] == ("Aluguel", 0, 2000)
    with_cancelled = reports.build_report(ENTRIES, "Ótica Teste", "Este mês", start, end, include_cancelled=True, today=TODAY)
    assert len(with_cancelled.rows) == 5 and with_cancelled.payable_open == 2000


def test_status_label_marks_overdue():
    assert reports.status_label(ENTRIES[2], TODAY) == "Atrasado"
    assert reports.status_label(ENTRIES[1], TODAY) == "Pendente"
    assert reports.status_label(ENTRIES[0], TODAY) == "Pago/Recebido"


def test_exports_create_valid_files(tmp_path):
    from openpyxl import load_workbook

    start, end = reports.period_bounds("Este mês", TODAY)
    report = reports.build_report(ENTRIES, "Ótica Teste", "Este mês", start, end, today=TODAY)
    xlsx = reports.export_xlsx(report, tmp_path / "r.xlsx")
    book = load_workbook(xlsx)
    assert book.sheetnames == ["Resumo", "Lançamentos", "Por categoria"]
    sheet = book["Lançamentos"]
    assert sheet.max_row == 1 + 4 + 1  # cabeçalho, 4 lançamentos, saldo
    assert sheet.cell(2, 11).value == -300  # saída com sinal negativo
    pdf = reports.export_pdf(report, tmp_path / "r.pdf")
    assert pdf.read_bytes()[:5] == b"%PDF-"
    empty = reports.build_report([], "Ótica Teste", "Hoje", TODAY, TODAY)
    assert reports.export_pdf(empty, tmp_path / "vazio.pdf").stat().st_size > 0
    assert reports.default_filename(report, "pdf") == "Relatorio_financeiro_Ótica_Teste_2026-10-01_a_2026-10-31.pdf"


@pytest.fixture(scope="module")
def qapp():
    from PySide6.QtWidgets import QApplication
    return QApplication.instance() or QApplication([])


def test_money_edit_types_through_cents(qapp):
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest

    from ui.components.money import MoneyEdit

    field = MoneyEdit()
    assert field.text() == "R$ 0,00"
    QTest.keyClicks(field, "1250")
    assert field.text() == "R$ 12,50" and field.value() == 12.5
    QTest.keyClick(field, Qt.Key.Key_Backspace)
    assert field.value() == 1.25
    field.setValue(1234.5)
    assert field.text() == "R$ 1.234,50"
    field._paste("2.345,67")
    assert field.value() == 2345.67


def test_money_edit_replaces_suggested_value(qapp):
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest

    from ui.components.money import MoneyEdit

    field = MoneyEdit()
    field.setValue(489.90)  # saldo sugerido pelo PDV
    QTest.keyClicks(field, "5000")
    assert field.value() == 50.0  # substitui, não soma dígitos ao saldo
    field.setValue(489.90)
    QTest.keyClick(field, Qt.Key.Key_Backspace)
    assert field.value() == 48.99  # Backspace corrige o valor existente
    QTest.keyClicks(field, "5")
    assert field.value() == 489.95
