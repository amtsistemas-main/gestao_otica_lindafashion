import pytest
from openpyxl import load_workbook

from modules.suppliers import spreadsheet

EXISTING = [
    {"id": "s1", "name": "Visão Lentes", "document": "12.345.678/0001-90", "city": "Maceió", "state": "AL", "active": True},
    {"id": "s2", "name": "Armações Sul", "document": "", "active": False},
]


def _fill(path, rows):
    book = load_workbook(path)
    sheet = book[spreadsheet.SHEET]
    sheet.delete_rows(2, sheet.max_row)
    for row in rows:
        sheet.append([row.get(key, "") for key in spreadsheet.COLUMNS])
    book.save(path)


def test_download_brings_current_suppliers(tmp_path):
    path = spreadsheet.create_spreadsheet(tmp_path / "f.xlsx", EXISTING)
    book = load_workbook(path)
    sheet = book[spreadsheet.SHEET]
    assert [cell.value for cell in sheet[1]][:2] == ["Nome de exibição", "Razão social"]
    filled = [row for row in sheet.iter_rows(min_row=2, values_only=True) if any(row)]
    assert len(filled) == 2
    assert "Instruções" in book.sheetnames
    records = spreadsheet.read_spreadsheet(path)
    assert [r["name"] for r in records] == ["Visão Lentes", "Armações Sul"]
    assert records[1]["active"] is False


def test_example_row_must_be_removed(tmp_path):
    path = spreadsheet.create_spreadsheet(tmp_path / "f.xlsx", [])
    with pytest.raises(ValueError, match="exemplo"):
        spreadsheet.read_spreadsheet(path)


def test_import_updates_by_document_or_name_and_creates_the_rest(tmp_path):
    path = spreadsheet.create_spreadsheet(tmp_path / "f.xlsx", [])
    _fill(path, [
        {"name": "Visão Lentes Atualizada", "document": "12345678000190", "city": "Recife", "state": "pe", "active": "Sim"},
        {"name": "armações sul", "phone": "(82) 3333-0000", "active": "Não"},
        {"name": "Nova Distribuidora", "document": "98.765.432/0001-10", "email": "vendas@nova.com.br"},
    ])
    records = spreadsheet.read_spreadsheet(path)
    assert records[0]["state"] == "PE"
    plan = spreadsheet.plan_import(records, EXISTING)
    assert [(sid, rec["name"]) for sid, rec in plan.to_update] == [("s1", "Visão Lentes Atualizada"), ("s2", "armações sul")]
    assert [rec["name"] for rec in plan.to_create] == ["Nova Distribuidora"]


def test_validation_reports_every_bad_row(tmp_path):
    path = spreadsheet.create_spreadsheet(tmp_path / "f.xlsx", [])
    _fill(path, [
        {"name": "", "city": "Maceió"},
        {"name": "Sem arroba", "email": "vendas.com"},
        {"name": "UF longa", "state": "Alagoas"},
        {"name": "Doc curto", "document": "123"},
    ])
    with pytest.raises(ValueError) as error:
        spreadsheet.read_spreadsheet(path)
    message = str(error.value)
    for expected in ("Linha 2", "obrigatório", "Linha 3", "e-mail", "Linha 4", "UF", "Linha 5", "CPF/CNPJ"):
        assert expected in message


def test_numeric_document_from_excel_keeps_leading_zeros():
    assert spreadsheet._text(1234567000190, "document") == "01234567000190"
    assert spreadsheet._text(12345678909.0, "document") == "12345678909"
    assert spreadsheet._text(57000000, "postal_code") == "57000000"


def test_repeated_supplier_in_sheet_is_rejected():
    records = [{"name": "A", "document": "12345678000190"}, {"name": "B", "document": "12.345.678/0001-90"}]
    with pytest.raises(ValueError, match="repetidos"):
        spreadsheet.plan_import(records, [])


def test_same_name_with_other_document_is_a_new_supplier():
    records = [{"name": "Visão Lentes", "document": "98765432000110"}]
    plan = spreadsheet.plan_import(records, EXISTING)
    assert not plan.to_update and len(plan.to_create) == 1
