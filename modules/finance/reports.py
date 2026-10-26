"""Finance periods and accountant reports (PDF and Excel)."""
from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any

PERIODS = ("Hoje", "Esta semana", "Este mês", "Este ano", "Últimos 30 dias", "Personalizado", "Todo o período")
REPORT_PERIODS = ("Hoje", "Esta semana", "Este mês", "Este ano", "Personalizado")
STATUS_LABELS = {"pendente": "Pendente", "pago": "Pago/Recebido", "cancelado": "Cancelado"}


def period_bounds(choice: str, today: date | None = None, custom_start: date | None = None,
                  custom_end: date | None = None) -> tuple[date | None, date | None]:
    """Whole calendar periods (the month runs to its last day, so upcoming bills are included).

    The week runs Sunday to Saturday, like the calendars of the system.
    """
    today = today or date.today()
    if choice == "Hoje":
        return today, today
    if choice == "Esta semana":
        start = today - timedelta(days=(today.weekday() + 1) % 7)
        return start, start + timedelta(days=6)
    if choice == "Este mês":
        next_month = (today.replace(day=28) + timedelta(days=4)).replace(day=1)
        return today.replace(day=1), next_month - timedelta(days=1)
    if choice == "Este ano":
        return today.replace(month=1, day=1), today.replace(month=12, day=31)
    if choice == "Últimos 30 dias":
        return today - timedelta(days=29), today
    if choice == "Personalizado" and custom_start and custom_end:
        return min(custom_start, custom_end), max(custom_start, custom_end)
    return None, None


def parse_day(value: Any) -> date | None:
    try:
        return date.fromisoformat(str(value or "")[:10])
    except ValueError:
        return None


def entry_day(entry: dict[str, Any]) -> date | None:
    """Paid entries count on the payment day; open and cancelled ones on the due day."""
    if entry.get("status") == "pago" and entry.get("paid_at"):
        return parse_day(entry.get("paid_at"))
    return parse_day(entry.get("due_date"))


def in_period(entry: dict[str, Any], start: date | None, end: date | None) -> bool:
    if start is None or end is None:
        return True
    day = entry_day(entry)
    return day is not None and start <= day <= end


def br_date(value: Any) -> str:
    day = parse_day(value)
    return day.strftime("%d/%m/%Y") if day else "—"


def money(value: Any) -> str:
    amount = float(value or 0)
    return ("R$ {:,.2f}".format(amount)).replace(",", "X").replace(".", ",").replace("X", ".")


def is_overdue(entry: dict[str, Any], today: date | None = None) -> bool:
    due = parse_day(entry.get("due_date"))
    return entry.get("status") == "pendente" and due is not None and due < (today or date.today())


def status_label(entry: dict[str, Any], today: date | None = None) -> str:
    if is_overdue(entry, today):
        return "Atrasado"
    return STATUS_LABELS.get(str(entry.get("status")), str(entry.get("status") or "—"))


@dataclass
class FinanceReport:
    store: str
    period_label: str
    start: date | None
    end: date | None
    rows: list[dict[str, Any]]
    income_paid: float = 0.0
    expense_paid: float = 0.0
    receivable_open: float = 0.0
    payable_open: float = 0.0
    overdue: float = 0.0
    cancelled: int = 0
    by_category: list[tuple[str, float, float]] = field(default_factory=list)
    generated_at: datetime = field(default_factory=datetime.now)

    @property
    def balance(self) -> float:
        return self.income_paid - self.expense_paid

    @property
    def range_text(self) -> str:
        if self.start is None or self.end is None:
            return "Todo o período"
        if self.start == self.end:
            return self.start.strftime("%d/%m/%Y")
        return f"{self.start.strftime('%d/%m/%Y')} a {self.end.strftime('%d/%m/%Y')}"


def build_report(entries: list[dict[str, Any]], store: str, period_label: str, start: date | None, end: date | None,
                 include_cancelled: bool = False, today: date | None = None) -> FinanceReport:
    selected = [entry for entry in entries if in_period(entry, start, end)]
    cancelled = sum(1 for entry in selected if entry.get("status") == "cancelado")
    if not include_cancelled:
        selected = [entry for entry in selected if entry.get("status") != "cancelado"]
    selected.sort(key=lambda entry: (entry_day(entry) or date.min, str(entry.get("description") or "")))
    report = FinanceReport(store=store, period_label=period_label, start=start, end=end, rows=selected, cancelled=cancelled)
    categories: dict[str, list[float]] = defaultdict(lambda: [0.0, 0.0])
    for entry in selected:
        amount = float(entry.get("amount") or 0)
        income = entry.get("direction") == "entrada"
        if entry.get("status") == "pago":
            if income:
                report.income_paid += amount
            else:
                report.expense_paid += amount
        elif entry.get("status") == "pendente":
            if income:
                report.receivable_open += amount
            else:
                report.payable_open += amount
            if is_overdue(entry, today):
                report.overdue += amount
        if entry.get("status") != "cancelado":
            categories[str(entry.get("category") or "Outros")][0 if income else 1] += amount
    report.by_category = sorted(((name, values[0], values[1]) for name, values in categories.items()),
                                key=lambda item: -(item[1] + item[2]))
    return report


def default_filename(report: FinanceReport, extension: str) -> str:
    stamp = report.start.strftime("%Y-%m-%d") if report.start else "completo"
    if report.end and report.end != report.start:
        stamp += f"_a_{report.end.strftime('%Y-%m-%d')}"
    store = "".join(ch for ch in report.store if ch.isalnum() or ch in " -_").strip().replace(" ", "_") or "Loja"
    return f"Relatorio_financeiro_{store}_{stamp}.{extension}"


# --------------------------------------------------------------------------- Excel
def export_xlsx(report: FinanceReport, path: str | Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    navy, blue, soft = "102957", "2165F5", "EAF0FC"
    money_format = '"R$" #,##0.00;[Red]-"R$" #,##0.00'
    thin = Side(style="thin", color="DCE3EE")
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor=navy)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumo"
    summary["A1"] = f"Relatório financeiro — {report.store}"
    summary["A1"].font = Font(bold=True, size=14, color=navy)
    summary["A2"] = f"Período: {report.period_label} ({report.range_text})"
    summary["A3"] = f"Gerado em {report.generated_at.strftime('%d/%m/%Y %H:%M')}"
    for cell in ("A2", "A3"):
        summary[cell].font = Font(color="4A5A75")
    lines = [
        ("Entradas recebidas", report.income_paid), ("Saídas pagas", report.expense_paid),
        ("Saldo realizado", report.balance), ("A receber (em aberto)", report.receivable_open),
        ("A pagar (em aberto)", report.payable_open), ("Em atraso", report.overdue),
    ]
    summary["A5"], summary["B5"] = "Indicador", "Valor"
    for cell in (summary["A5"], summary["B5"]):
        cell.font, cell.fill = header_font, header_fill
    for offset, (label, value) in enumerate(lines, start=6):
        summary.cell(offset, 1, label)
        amount = summary.cell(offset, 2, round(value, 2))
        amount.number_format = money_format
        if label == "Saldo realizado":
            summary.cell(offset, 1).font = Font(bold=True)
            amount.font = Font(bold=True)
    summary.cell(13, 1, f"Lançamentos no período: {len(report.rows)}")
    if report.cancelled:
        summary.cell(14, 1, f"Cancelados no período: {report.cancelled}")
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 18

    sheet = workbook.create_sheet("Lançamentos")
    headers = ["Data", "Vencimento", "Pagamento", "Tipo", "Descrição", "Categoria", "Cliente / fornecedor",
               "Forma de pagamento", "Situação", "Referência", "Valor"]
    sheet.append(headers)
    for column in range(1, len(headers) + 1):
        cell = sheet.cell(1, column)
        cell.font, cell.fill = header_font, header_fill
        cell.alignment = Alignment(vertical="center")
    for entry in report.rows:
        income = entry.get("direction") == "entrada"
        amount = float(entry.get("amount") or 0)
        sheet.append([
            entry_day(entry), parse_day(entry.get("due_date")), parse_day(entry.get("paid_at")),
            "Entrada" if income else "Saída", str(entry.get("description") or ""), str(entry.get("category") or ""),
            str(entry.get("counterparty") or ""), str(entry.get("payment_method") or ""), status_label(entry),
            str(entry.get("reference") or ""), round(amount if income else -amount, 2),
        ])
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        for cell in row[:3]:
            cell.number_format = "DD/MM/YYYY"
        row[-1].number_format = money_format
        for cell in row:
            cell.border = Border(bottom=thin)
    total_row = sheet.max_row + 1
    sheet.cell(total_row, len(headers) - 1, "Saldo do período").font = Font(bold=True)
    total = sheet.cell(total_row, len(headers), f"=SUM({get_column_letter(len(headers))}2:{get_column_letter(len(headers))}{total_row - 1})"
                       if report.rows else 0)
    total.number_format, total.font = money_format, Font(bold=True, color=blue)
    widths = [12, 12, 12, 9, 42, 18, 28, 20, 15, 16, 15]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    if report.rows:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{total_row - 1}"

    categories = workbook.create_sheet("Por categoria")
    categories.append(["Categoria", "Entradas", "Saídas", "Resultado"])
    for column in range(1, 5):
        cell = categories.cell(1, column)
        cell.font, cell.fill = header_font, header_fill
    for name, income, expense in report.by_category:
        categories.append([name, round(income, 2), round(expense, 2), round(income - expense, 2)])
    for row in categories.iter_rows(min_row=2, max_row=categories.max_row):
        for cell in row[1:]:
            cell.number_format = money_format
    for letter, width in zip("ABCD", (28, 16, 16, 16)):
        categories.column_dimensions[letter].width = width
    for sheet_ in (summary, categories):
        sheet_.sheet_view.showGridLines = False
    summary.cell(5, 1).fill = summary.cell(5, 2).fill = PatternFill("solid", fgColor=navy)
    summary.cell(8, 1).fill = summary.cell(8, 2).fill = PatternFill("solid", fgColor=soft)

    path = Path(path)
    workbook.save(path)
    return path


# --------------------------------------------------------------------------- PDF
def export_pdf(report: FinanceReport, path: str | Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from modules.sales.document_pdf import _register_fonts

    regular, bold = _register_fonts()
    navy, blue, ink, slate, line, soft = (colors.HexColor(c) for c in ("#102957", "#2165F5", "#1E2B42", "#4A5A75", "#DCE3EE", "#F5F7FB"))
    green, red = colors.HexColor("#0E8A5B"), colors.HexColor("#C03030")
    page = landscape(A4)
    margin = 12 * mm
    width = page[0] - 2 * margin
    cell = ParagraphStyle("cell", fontName=regular, fontSize=7.6, leading=9.4, textColor=ink)
    cell_right = ParagraphStyle("cellRight", parent=cell, alignment=2)
    head = ParagraphStyle("head", fontName=bold, fontSize=7.6, leading=9.4, textColor=colors.white)
    head_right = ParagraphStyle("headRight", parent=head, alignment=2)
    section = ParagraphStyle("section", fontName=bold, fontSize=10.5, leading=13, textColor=navy, spaceBefore=6, spaceAfter=5,
                             keepWithNext=1)

    def esc(text: Any) -> str:
        return str(text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    def decorate(canvas, doc) -> None:
        canvas.saveState()
        canvas.setFillColor(navy)
        canvas.rect(0, page[1] - 22 * mm, page[0], 22 * mm, stroke=0, fill=1)
        canvas.setFillColor(colors.white)
        canvas.setFont(bold, 14)
        canvas.drawString(margin, page[1] - 11 * mm, "Relatório financeiro")
        canvas.setFont(regular, 9)
        canvas.drawString(margin, page[1] - 17 * mm, f"{report.store}  •  {report.period_label}: {report.range_text}")
        canvas.drawRightString(page[0] - margin, page[1] - 11 * mm, f"Gerado em {report.generated_at.strftime('%d/%m/%Y %H:%M')}")
        canvas.setFillColor(slate)
        canvas.setFont(regular, 7.5)
        canvas.drawString(margin, 7 * mm, "Valores de entradas e saídas pagas consideram a data do pagamento; em aberto, a data de vencimento.")
        canvas.drawRightString(page[0] - margin, 7 * mm, f"Página {doc.page}")
        canvas.restoreState()

    story: list[Any] = []
    cards = [("Entradas recebidas", report.income_paid, green), ("Saídas pagas", report.expense_paid, red),
             ("Saldo realizado", report.balance, blue), ("A receber", report.receivable_open, ink),
             ("A pagar", report.payable_open, ink), ("Em atraso", report.overdue, red)]
    card_cells = [[Paragraph(f"<font size=7.5 color='#4A5A75'>{label}</font><br/><font name='{bold}' size=12 color='{tone.hexval().replace('0x', '#')}'>{money(value)}</font>",
                             ParagraphStyle("card", fontName=regular, leading=15)) for label, value, tone in cards]]
    card_table = Table(card_cells, colWidths=[width / 6] * 6)
    card_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), soft), ("BOX", (0, 0), (-1, -1), 0.6, line),
        ("LINEAFTER", (0, 0), (-2, -1), 0.6, line), ("TOPPADDING", (0, 0), (-1, -1), 7),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7), ("LEFTPADDING", (0, 0), (-1, -1), 8),
    ]))
    story += [card_table, Spacer(1, 8)]

    story.append(Paragraph(f"Lançamentos ({len(report.rows)})", section))
    columns = [("Data", 17), ("Vencimento", 19), ("Tipo", 13), ("Descrição", 66), ("Categoria", 28),
               ("Cliente / fornecedor", 44), ("Forma", 26), ("Situação", 22), ("Valor", 24)]
    scale = width / (sum(w for _n, w in columns) * mm)
    widths = [w * mm * scale for _n, w in columns]
    data = [[Paragraph(name, head_right if name == "Valor" else head) for name, _w in columns]]
    for entry in report.rows:
        income = entry.get("direction") == "entrada"
        amount = float(entry.get("amount") or 0)
        tone = "#0E8A5B" if income else "#C03030"
        data.append([
            Paragraph(br_date(entry_day(entry).isoformat() if entry_day(entry) else None), cell),
            Paragraph(br_date(entry.get("due_date")), cell),
            Paragraph("Entrada" if income else "Saída", cell),
            Paragraph(esc(entry.get("description")), cell), Paragraph(esc(entry.get("category")), cell),
            Paragraph(esc(entry.get("counterparty")) or "—", cell), Paragraph(esc(entry.get("payment_method")) or "—", cell),
            Paragraph(status_label(entry), cell),
            Paragraph(f"<font color='{tone}'>{'' if income else '−'}{money(amount)}</font>", cell_right),
        ])
    if not report.rows:
        data.append([Paragraph("Nenhum lançamento no período.", cell)] + [""] * (len(columns) - 1))
    entries_table = Table(data, colWidths=widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), navy), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, 1), (-1, -1), 0.4, line), ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4), ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
    ]
    style += [("BACKGROUND", (0, row), (-1, row), soft) for row in range(2, len(data), 2)]
    if not report.rows:
        style.append(("SPAN", (0, 1), (-1, 1)))
    entries_table.setStyle(TableStyle(style))
    story += [entries_table, Spacer(1, 10)]

    if report.by_category:
        story.append(Paragraph("Resultado por categoria", section))
        rows = [[Paragraph(name, head if name == "Categoria" else head_right) for name in ("Categoria", "Entradas", "Saídas", "Resultado")]]
        for name, income, expense in report.by_category:
            rows.append([Paragraph(esc(name), cell), Paragraph(money(income), cell_right),
                         Paragraph(money(expense), cell_right), Paragraph(f"<b>{money(income - expense)}</b>", cell_right)])
        category_table = Table(rows, colWidths=[width * 0.4, width * 0.2, width * 0.2, width * 0.2], repeatRows=1)
        category_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), navy), ("LINEBELOW", (0, 1), (-1, -1), 0.4, line),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(category_table)

    path = Path(path)
    doc = SimpleDocTemplate(str(path), pagesize=page, leftMargin=margin, rightMargin=margin,
                            topMargin=27 * mm, bottomMargin=14 * mm, title=f"Relatório financeiro — {report.store}",
                            author="AMT Sistemas & Soluções")
    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return path
