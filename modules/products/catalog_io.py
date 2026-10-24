"""XLS/XLSX product import/template and PDF catalog export helpers."""
from __future__ import annotations

from datetime import datetime
from decimal import Decimal, InvalidOperation
from html import escape
from io import BytesIO
from urllib.parse import urlparse
from urllib.request import urlopen
from pathlib import Path
from typing import Any, Iterable

from modules.products.models import Product

HEADERS = [
    "Código de barras", "SKU", "Nome", "Categoria", "Marca", "Preço de custo",
    "Preço de venda", "Desconto (%)", "Estoque", "Estoque mínimo", "Descrição",
    "URL da imagem", "Ativo",
]
FIELDS = [
    "barcode", "sku", "name", "category", "brand", "cost_price", "sale_price",
    "discount_percent", "stock_quantity", "minimum_stock", "description", "image_url", "active",
]


def _decimal(value: Any, field: str) -> Decimal:
    if value in (None, ""):
        return Decimal("0")
    raw = str(value).strip()
    if isinstance(value, str) and "," in raw:
        raw = raw.replace(".", "").replace(",", ".") if raw.rfind(",") > raw.rfind(".") else raw.replace(",", "")
    try:
        return Decimal(raw)
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"valor inválido no campo {field}: {value!r}") from exc


def _as_text(value: Any) -> str:
    return "" if value is None else str(value).strip()


def _as_int(value: Any, field: str) -> int:
    parsed = _decimal(value, field)
    if parsed != parsed.to_integral_value():
        raise ValueError(f"{field} deve ser um número inteiro.")
    return int(parsed)


def _as_bool(value: Any) -> bool:
    if value is None or value == "":
        return True
    if isinstance(value, bool):
        return value
    return str(value).strip().casefold() in {"1", "sim", "s", "true", "ativo", "yes", "y"}


def create_template(path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Produtos"
    sheet.append(HEADERS)
    sheet.append(["7890000000012", "ARM-001", "Armação demonstrativa", "Armações", "Marca exemplo",
                  80.00, 159.90, 0, 10, 2, "Remova esta linha de exemplo antes de importar", "", "Sim"])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}1"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="102957")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30
    widths = [20, 16, 30, 20, 20, 16, 16, 14, 12, 16, 38, 36, 12]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    notes = workbook.create_sheet("Instruções")
    notes.append(["Como importar o catálogo"])
    notes.append(["Preencha a aba Produtos; a primeira linha contém os cabeçalhos obrigatórios."])
    notes.append(["Remova a linha de exemplo antes de importar. Nome é obrigatório."])
    notes.append(["Preços aceitam número decimal; use a coluna Desconto (%) entre 0 e 100."])
    notes.append(["Estoque e Estoque mínimo devem ser inteiros não negativos."])
    notes.append(["Ativo: Sim ou Não. URL da imagem é opcional."])
    notes.append(["O sistema aceita .xlsx e arquivos .xls antigos quando a dependência xlrd estiver instalada."])
    notes.column_dimensions["A"].width = 110
    workbook.save(path)


def read_products(path: str) -> list[Product]:
    suffix = Path(path).suffix.lower()
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook["Produtos"] if "Produtos" in workbook.sheetnames else workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = [_as_text(item).casefold() for item in next(rows)]
        except StopIteration:
            return []
        index_by_header = {header: index for index, header in enumerate(headers)}
        mapping = {
            "barcode": ["código de barras", "codigo de barras", "barcode"],
            "sku": ["sku", "código interno", "codigo interno"],
            "name": ["nome", "produto", "name"],
            "category": ["categoria", "category"],
            "brand": ["marca", "brand"],
            "cost_price": ["preço de custo", "preco de custo", "cost_price"],
            "sale_price": ["preço de venda", "preco de venda", "sale_price"],
            "discount_percent": ["desconto (%)", "desconto", "discount_percent"],
            "stock_quantity": ["estoque", "quantidade em estoque", "stock_quantity"],
            "minimum_stock": ["estoque mínimo", "estoque minimo", "minimum_stock"],
            "description": ["descrição", "descricao", "description"],
            "image_url": ["url da imagem", "imagem", "image_url"],
            "active": ["ativo", "active"],
        }
        def get_value(row: tuple[Any, ...], key: str) -> Any:
            for name in mapping[key]:
                if name in index_by_header:
                    idx = index_by_header[name]
                    return row[idx] if idx < len(row) else None
            return None
        data_rows = list(rows)
        workbook.close()
    elif suffix == ".xls":
        import xlrd
        workbook = xlrd.open_workbook(path, on_demand=True)
        sheet = workbook.sheet_by_name("Produtos") if "Produtos" in workbook.sheet_names() else workbook.sheet_by_index(0)
        headers = [_as_text(sheet.cell_value(0, col)).casefold() for col in range(sheet.ncols)] if sheet.nrows else []
        index_by_header = {header: index for index, header in enumerate(headers)}
        mapping = {
            "barcode": ["código de barras", "codigo de barras", "barcode"],
            "sku": ["sku", "código interno", "codigo interno"],
            "name": ["nome", "produto", "name"],
            "category": ["categoria", "category"], "brand": ["marca", "brand"],
            "cost_price": ["preço de custo", "preco de custo", "cost_price"],
            "sale_price": ["preço de venda", "preco de venda", "sale_price"],
            "discount_percent": ["desconto (%)", "desconto", "discount_percent"],
            "stock_quantity": ["estoque", "quantidade em estoque", "stock_quantity"],
            "minimum_stock": ["estoque mínimo", "estoque minimo", "minimum_stock"],
            "description": ["descrição", "descricao", "description"],
            "image_url": ["url da imagem", "imagem", "image_url"], "active": ["ativo", "active"],
        }
        data_rows = [tuple(sheet.cell_value(row, col) for col in range(sheet.ncols)) for row in range(1, sheet.nrows)]
        def get_value(row: tuple[Any, ...], key: str) -> Any:
            for name in mapping[key]:
                if name in index_by_header:
                    idx = index_by_header[name]
                    return row[idx] if idx < len(row) else None
            return None
        workbook.release_resources()
    else:
        raise ValueError("Selecione uma planilha .xlsx ou .xls.")

    products: list[Product] = []
    errors: list[str] = []
    for row_number, row in enumerate(data_rows, start=2):
        if not any(value not in (None, "") for value in row):
            continue
        try:
            product = Product(
                barcode=_as_text(get_value(row, "barcode")),
                sku=_as_text(get_value(row, "sku")),
                name=_as_text(get_value(row, "name")),
                category=_as_text(get_value(row, "category")),
                brand=_as_text(get_value(row, "brand")),
                cost_price=_decimal(get_value(row, "cost_price"), "Preço de custo"),
                sale_price=_decimal(get_value(row, "sale_price"), "Preço de venda"),
                discount_percent=_decimal(get_value(row, "discount_percent"), "Desconto"),
                stock_quantity=_as_int(get_value(row, "stock_quantity"), "Estoque"),
                minimum_stock=_as_int(get_value(row, "minimum_stock"), "Estoque mínimo"),
                description=_as_text(get_value(row, "description")),
                image_url=_as_text(get_value(row, "image_url")),
                active=_as_bool(get_value(row, "active")),
            )
            product.validate()
            products.append(product)
        except Exception as exc:
            errors.append(f"Linha {row_number}: {exc}")
    if errors:
        raise ValueError("A planilha contém dados inválidos; nenhum produto foi importado:\n" + "\n".join(errors[:30]))
    if not products:
        raise ValueError("A planilha não contém produtos para importar.")
    return products


def export_pdf(path: str, products: Iterable[Product]) -> None:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    items = list(products)
    document = SimpleDocTemplate(
        path, pagesize=landscape(A4), rightMargin=12 * mm, leftMargin=12 * mm,
        topMargin=12 * mm, bottomMargin=12 * mm, title="Catálogo de Produtos - Ótica Linda Fashion",
        author="Ótica Linda Fashion",
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("CatalogTitle", parent=styles["Title"], textColor=colors.HexColor("#102957"), fontName="Helvetica-Bold", fontSize=20, leading=24, alignment=TA_LEFT, spaceAfter=3*mm)
    sub_style = ParagraphStyle("CatalogSubtitle", parent=styles["Normal"], textColor=colors.HexColor("#60799E"), fontSize=9, leading=12)
    cell_style = ParagraphStyle("CatalogCell", parent=styles["BodyText"], fontSize=7.5, leading=9)
    head_style = ParagraphStyle("CatalogHead", parent=cell_style, textColor=colors.white, fontName="Helvetica-Bold")
    story = [
        Paragraph("ÓTICA LINDA FASHION", title_style),
        Paragraph("CATÁLOGO DE PRODUTOS · Gerado em " + datetime.now().strftime("%d/%m/%Y %H:%M"), sub_style),
        Spacer(1, 6*mm),
    ]
    rows = [[Paragraph(text, head_style) for text in ["Imagem", "Código / SKU", "Produto", "Categoria", "Marca", "Custo", "Venda", "Desc.", "Preço final", "Estoque"]]]
    for product in items:
        image_cell: Any = "—"
        if product.image_url:
            try:
                parsed = urlparse(product.image_url)
                if parsed.scheme == "file":
                    image_bytes = Path(parsed.path).read_bytes()
                elif parsed.scheme in {"http", "https"}:
                    with urlopen(product.image_url, timeout=3) as response:
                        image_bytes = response.read(8 * 1024 * 1024 + 1)
                    if len(image_bytes) > 8 * 1024 * 1024:
                        image_bytes = b""
                else:
                    image_bytes = b""
                if image_bytes:
                    image_cell = Image(BytesIO(image_bytes), width=15*mm, height=15*mm, kind="proportional")
            except Exception:
                image_cell = "—"
        rows.append([
            image_cell,
            Paragraph(escape(product.barcode or product.sku or "—"), cell_style),
            Paragraph(escape(product.name), cell_style),
            Paragraph(escape(product.category or "—"), cell_style),
            Paragraph(escape(product.brand or "—"), cell_style),
            _currency_text(product.cost_price), _currency_text(product.sale_price),
            f"{product.discount_percent:.2f}%".replace(".", ","),
            _currency_text(product.final_price), str(product.stock_quantity),
        ])
    table = Table(rows, repeatRows=1, colWidths=[19*mm, 26*mm, 48*mm, 28*mm, 24*mm, 21*mm, 21*mm, 15*mm, 24*mm, 17*mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#102957")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#E1E8F2")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F5F8FD")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (5, 1), (-1, -1), "RIGHT"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(table)
    story.append(Spacer(1, 4*mm))
    story.append(Paragraph(f"Total de produtos no catálogo: {len(items)}", sub_style))
    document.build(story)


def _currency_text(value: Decimal) -> str:
    return f"R$ {value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
