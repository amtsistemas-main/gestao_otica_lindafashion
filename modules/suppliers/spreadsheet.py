"""Supplier spreadsheet: download (current suppliers or an example), fill in and import.

Import updates suppliers matched by CPF/CNPJ (or, without a document, by display name)
and creates the others. Every row is validated before anything is written.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from modules.suppliers.models import FIELD_LABELS, SUPPLIER_FIELDS

SHEET = "Fornecedores"
COLUMNS = [key for key in SUPPLIER_FIELDS if key != "active"] + ["active"]
HEADERS = [FIELD_LABELS.get(key, "Ativo").replace(" *", "") for key in COLUMNS]
HEADERS[0] = "Nome de exibição"
REQUIRED_HEADER = "nome de exibição"
EXAMPLE = {
    "name": "Distribuidora Exemplo", "legal_name": "Distribuidora Exemplo Ltda", "trade_name": "Exemplo Óptica",
    "document": "12.345.678/0001-90", "contact_name": "Ana Souza", "phone": "(82) 3333-4444",
    "whatsapp": "(82) 99999-0000", "email": "contato@exemplo.com.br", "postal_code": "57000-000",
    "street": "Rua das Lentes", "address_number": "100", "neighborhood": "Centro", "city": "Maceió", "state": "AL",
    "payment_terms": "30/60 dias", "notes": "APAGUE ESTA LINHA DE EXEMPLO ANTES DE IMPORTAR",
}
_TEXT_COLUMNS = {"document", "phone", "whatsapp", "postal_code", "address_number", "state_registration",
                 "municipal_registration", "bank_branch", "bank_account", "pix_key"}


@dataclass
class ImportPlan:
    to_create: list[dict[str, Any]] = field(default_factory=list)
    to_update: list[tuple[str, dict[str, Any]]] = field(default_factory=list)


def digits(value: Any) -> str:
    return re.sub(r"\D", "", str(value or ""))


def _text(value: Any, key: str) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    # Excel turns CPF/CNPJ/CEP typed as numbers into integers and drops leading zeros.
    if isinstance(value, int) and key == "document":
        text = text.zfill(11 if len(text) <= 11 else 14)
    elif isinstance(value, int) and key == "postal_code":
        text = text.zfill(8)
    return text


def _bool(value: Any) -> bool:
    if value is None or str(value).strip() == "":
        return True
    if isinstance(value, bool):
        return value
    return str(value).strip().casefold() in {"1", "sim", "s", "true", "ativo", "yes", "y"}


def create_spreadsheet(path: str | Path, suppliers: list[dict[str, Any]]) -> Path:
    """Current suppliers ready to edit; with none registered, a single example row."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET
    sheet.append(HEADERS)
    rows = suppliers or [EXAMPLE]
    for supplier in rows:
        values = []
        for key in COLUMNS:
            if key == "active":
                values.append("Sim" if supplier.get("active", True) else "Não")
            else:
                values.append(str(supplier.get(key) or ""))
        sheet.append(values)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="102957")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.cell(1, 1).fill = PatternFill("solid", fgColor="2165F5")
    sheet.row_dimensions[1].height = 32
    for index, key in enumerate(COLUMNS, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = 34 if key in ("name", "legal_name", "notes", "street") else 20
        if key in _TEXT_COLUMNS:
            for row in range(2, max(len(rows) + 1, 500) + 1):
                sheet[f"{letter}{row}"].number_format = "@"
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows) + 1}"

    notes = workbook.create_sheet("Instruções")
    for line in (
        "Como cadastrar ou atualizar fornecedores pela planilha",
        "",
        "1. Preencha a aba Fornecedores: uma linha por fornecedor. A coluna Nome de exibição (em azul) é obrigatória.",
        "2. Fornecedores já cadastrados são reconhecidos pelo CPF/CNPJ (ou, sem documento, pelo nome de exibição) e ATUALIZADOS.",
        "3. Linhas que não correspondem a nenhum fornecedor são CRIADAS como novos cadastros.",
        "4. Colunas deixadas em branco num fornecedor existente apagam aquela informação; para manter, deixe o valor atual.",
        "5. UF com 2 letras (ex.: AL). E-mail no formato nome@dominio. CPF com 11 e CNPJ com 14 números.",
        "6. Ativo: Sim ou Não (em branco = Sim).",
        "7. Se houver linha de exemplo, apague-a antes de importar.",
        "8. Em Fornecedores, clique em Importar planilha: o sistema confere tudo e mostra um resumo antes de gravar.",
    ):
        notes.append([line])
    notes["A1"].font = Font(bold=True, size=13, color="102957")
    notes.column_dimensions["A"].width = 120
    path = Path(path)
    workbook.save(path)
    return path


def read_spreadsheet(path: str | Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    if Path(path).suffix.lower() != ".xlsx":
        raise ValueError("Selecione uma planilha .xlsx (baixe o modelo pelo botão Baixar planilha).")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[SHEET] if SHEET in workbook.sheetnames else workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            header_row = [str(value or "").strip().casefold().replace(" *", "") for value in next(rows)]
        except StopIteration:
            raise ValueError("A planilha está vazia.") from None
        labels = {header.casefold(): key for header, key in zip(HEADERS, COLUMNS)}
        index = {labels[header]: position for position, header in enumerate(header_row) if header in labels}
        if "name" not in index:
            raise ValueError("Coluna “Nome de exibição” não encontrada. Use a planilha baixada pelo sistema.")
        records: list[dict[str, Any]] = []
        errors: list[str] = []
        for number, row in enumerate(rows, start=2):
            if not any(value not in (None, "") for value in row):
                continue
            record = {}
            for key, position in index.items():
                value = row[position] if position < len(row) else None
                record[key] = _bool(value) if key == "active" else _text(value, key)
            record.setdefault("active", True)
            problems = _validate(record)
            if problems:
                errors.append(f"Linha {number}: " + "; ".join(problems))
            else:
                records.append(record)
    finally:
        workbook.close()
    if errors:
        raise ValueError("A planilha tem dados a corrigir (nada foi gravado):\n" + "\n".join(errors[:30])
                         + (f"\n… e mais {len(errors) - 30} linha(s)." if len(errors) > 30 else ""))
    if not records:
        raise ValueError("A planilha não tem fornecedores para importar.")
    return records


def _validate(record: dict[str, Any]) -> list[str]:
    problems = []
    if not record.get("name"):
        problems.append("Nome de exibição é obrigatório")
    if "APAGUE ESTA LINHA" in str(record.get("notes") or "").upper():
        problems.append("linha de exemplo — apague antes de importar")
    email = record.get("email") or ""
    if email and ("@" not in email or "." not in email.rsplit("@", 1)[-1]):
        problems.append(f"e-mail inválido ({email})")
    state = record.get("state") or ""
    if state and not re.fullmatch(r"[A-Za-z]{2}", state):
        problems.append(f"UF deve ter 2 letras ({state})")
    elif state:
        record["state"] = state.upper()
    document = digits(record.get("document"))
    if record.get("document") and len(document) not in (11, 14):
        problems.append(f"CPF/CNPJ deve ter 11 ou 14 números ({record.get('document')})")
    return problems


def _key(record: dict[str, Any]) -> tuple[str, str]:
    document = digits(record.get("document"))
    return ("doc", document) if document else ("name", str(record.get("name") or "").strip().casefold())


def plan_import(records: list[dict[str, Any]], existing: list[dict[str, Any]]) -> ImportPlan:
    """Matches by CPF/CNPJ, then by display name; repeated rows in the sheet are rejected."""
    by_document = {digits(row.get("document")): row for row in existing if digits(row.get("document"))}
    by_name = {str(row.get("name") or "").strip().casefold(): row for row in existing}
    seen: dict[tuple[str, str], int] = {}
    duplicates = []
    plan = ImportPlan()
    for position, record in enumerate(records, start=2):
        key = _key(record)
        if key in seen:
            duplicates.append(f"“{record['name']}” aparece mais de uma vez na planilha")
            continue
        seen[key] = position
        match = by_document.get(key[1]) if key[0] == "doc" else None
        if match is None:
            match = by_name.get(str(record.get("name") or "").strip().casefold())
            # A different document under the same name is a different supplier.
            if match is not None and key[0] == "doc" and digits(match.get("document")) not in ("", key[1]):
                match = None
        if match is not None:
            plan.to_update.append((str(match["id"]), record))
        else:
            plan.to_create.append(record)
    if duplicates:
        raise ValueError("A planilha tem fornecedores repetidos (nada foi gravado):\n" + "\n".join(duplicates[:30]))
    return plan
